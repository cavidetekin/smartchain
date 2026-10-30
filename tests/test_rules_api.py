"""Koşullu biçim kuralları (U-24, K-95, K-96): sunucuda değerlendirme, alan karşılaştırması, "Her zaman", görünüm, Excel.

Sentetik v0 ve yerel SQL Server. Veriyi değiştiren testler bir işlem içinde çalışır ve işlem geri alınır.
"""
import io
from datetime import date, datetime, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.lists import MATERIALS
from app.api.lists.definition import ListDefinition
from app.api.lists.request import ExportRequest, QueryRequest
from app.api.services import export as export_service, lists as list_service

from test_export_api import user  # noqa: F401  (iz satırlarını temizleyen deneme kullanıcısı)
from test_lists_api import all_rows, api, query  # noqa: F401
from test_views_api import create, listed, users  # noqa: F401

pytestmark = pytest.mark.db

FIELDS = ["Code", "Name", "MaterialGroup", "SafetyStock", "MaxStock", "UpdatedAt", "IsSellable"]


def hits(rows: list[dict], i: int = 0) -> list[str]:
    return [r["Code"] for r in rows if i in r["_rules"]]


def codes(rows: list[dict]) -> list[str]:
    return [r["Code"] for r in rows]


def rule(field: str, **condition) -> dict:
    return {"field": field, "condition": condition}


# --- filtreyle aynı tanım: aynı koşul filtre olarak verilince aynı satırlar ---------------------------------

SAME_AS_FILTER = [
    rule("Name", op="contains", value="mil"),
    rule("Name", op="contains", value="MİL"),          # Türkçe İ/i, büyük/küçük harf duyarsız
    rule("Name", op="equals", value="havalandırma tapası"),
    rule("Code", op="starts_with", value="sat-0"),
    rule("Name", op="empty"),
    rule("AdditionalText1", op="empty"),
    rule("SafetyStock", op="gt", value=10),
    rule("SafetyStock", op="lt", value=50),
    rule("SafetyStock", op="eq", value=0),
    rule("SafetyStock", op="between", min=10, max=200),
    rule("SafetyStock", op="empty"),
    rule("UpdatedAt", op="relative", value="this_week"),
    rule("IsSellable", op="in", values=[True]),
]


@pytest.mark.parametrize("r", SAME_AS_FILTER, ids=lambda r: f"{r['field']}-{r['condition']['op']}")
def test_rule_marks_same_rows_as_filter(api, r):
    fields = list(dict.fromkeys(FIELDS + [r["field"]]))
    rows = all_rows(api, fields=fields, rules=[r])
    flt = {"field": r["field"], **r["condition"]}
    expected = codes(all_rows(api, fields=["Code"], filters=[flt]))
    assert sorted(hits(rows)) == sorted(expected)


def test_text_rules_on_list_field_and_in(api):
    """U-24: liste filtreli metin alanında da eşittir, içerir, ile başlar; "listedekilerden biri" harf duyarsız."""
    rules = [rule("MaterialGroup", op="starts_with", value="çub"), rule("MaterialGroup", op="in", values=["ÇUBUK", "levha"]),
             rule("MaterialGroup", op="equals", value="çubuk")]
    rows = all_rows(api, fields=["Code", "MaterialGroup"], rules=rules)
    group = {r["Code"]: (r["MaterialGroup"] or "") for r in rows}
    assert sorted(hits(rows, 0)) == sorted(c for c, g in group.items() if g.startswith("Çub"))
    assert sorted(hits(rows, 1)) == sorted(c for c, g in group.items() if g in ("Çubuk", "Levha"))
    assert sorted(hits(rows, 2)) == sorted(hits(rows, 0)) and hits(rows, 1)


@pytest.mark.parametrize("field", ["SafetyStock", "UpdatedAt", "Name", "MaterialGroup", "IsSellable", "AdditionalText1"])
def test_always_rule_marks_every_row_including_empty(api, field):
    """"Her zaman": koşulsuz kolon biçimi, her tipte; boş hücreler dahil bütün satırlarda tutar."""
    rows = all_rows(api, fields=["Code", field], rules=[rule(field, op="always")])
    assert len(rows) == 200 and all(r["_rules"] == [0] for r in rows)
    if field == "SafetyStock":
        assert any(r[field] in (None, "") for r in rows)  # boş hücre de var ve o da işaretli


def test_rows_without_rules_have_no_rules_key(api):
    assert "_rules" not in query(api, fields=["Code"], limit=1)["rows"][0]


# --- doğrulama ---------------------------------------------------------------------------------------------

@pytest.mark.parametrize("rules, field, problem", [
    ([rule("Name", op="gt", value=1)], "rules.0.condition.op", "kullanılamaz"),
    ([rule("Nope", op="empty")], "rules.0.field", "Bilinmeyen alan"),
    ([rule("SafetyStock", op="lt", field_ref="Name")], "rules.0.condition.field_ref", "tipi uymuyor"),
    ([rule("SafetyStock", op="lt", field_ref="Nope")], "rules.0.condition.field_ref", "Bilinmeyen alan"),
    ([rule("SafetyStock", op="eq", field_ref="SafetyStock")], "rules.0.condition.field_ref", "kendisiyle"),
    ([rule("SafetyStock", op="gt", field_ref="MaxStock", value=3)], "rules.0.condition.field_ref", "birlikte"),
    ([rule("SafetyStock", op="between", field_ref="MaxStock", min=1)], "rules.0.condition.field_ref", "yalnız sayı"),
    ([rule("UpdatedAt", op="gt", value="dün")], "rules.0.condition.value", "Tarih"),
    ([rule("IsSellable", op="in", values=["evet"])], "rules.0.condition.values", "true ya da false"),
    ([rule("Name", op="contains", value="  ")], "rules.0.condition.value", "Metin"),
    ([rule("UpdatedAt", op="gt", field_ref="UpdatedAt", empty_as_zero=True)], "rules.0.condition.empty_as_zero", "yalnız sayı"),
    ([rule("SafetyStock", op="gt", value=1, empty_as_zero=False)], "rules.0.condition.empty_as_zero", "yalnız sayı"),
])
def test_invalid_rules_are_422(api, rules, field, problem):
    r = api.post("/api/materials/query", json={"rules": rules})
    assert r.status_code == 422, r.text
    details = r.json()["error"]["details"]
    assert any(d["field"] == field and problem in d["problem"] for d in details), details


def test_rule_limit_and_format_validation(api):
    r = api.post("/api/materials/query", json={"rules": [rule("Name", op="empty")] * 21})
    assert r.status_code == 422 and r.json()["error"]["details"][0]["field"] == "body.rules"


# --- alan karşılaştırması (K-95): işlem içinde veri kurulur, sonra geri alınır ------------------------------

@pytest.fixture
def txn(db_engine):
    conn = db_engine.connect()
    trans = conn.begin()
    session = Session(bind=conn)
    try:
        yield conn, session
    finally:
        session.close()
        trans.rollback()
        conn.close()


def test_number_field_ref(txn):
    conn, session = txn
    ids = conn.execute(text("SELECT TOP 6 Id FROM core.Material WHERE VersionTo IS NULL ORDER BY Code")).scalars().all()
    for i, (safety, mx) in zip(ids, [(5, 10), (10, 10), (20, 10), (None, 10), (5, None), (0, 0.5)]):
        conn.execute(text("UPDATE core.Material SET SafetyStock = :s, MaxStock = :m WHERE Id = :i"),
                     {"s": safety, "m": mx, "i": i})
    ops = ("lt", "gt", "eq")
    req = QueryRequest(fields=["Code", "SafetyStock", "MaxStock"], limit=100, rules=[  # varsayılan: boşları 0 say
        *[rule("SafetyStock", op=op, field_ref="MaxStock") for op in ops],
        *[rule("SafetyStock", op=op, field_ref="MaxStock", empty_as_zero=False) for op in ops]])
    rows = list_service.query(session, MATERIALS, req)["rows"]

    def expect(op, zero):
        out = []
        for r in rows:
            a, b = r["SafetyStock"], r["MaxStock"]
            if zero:
                a, b = a or 0, b or 0
            elif a is None or b is None:
                continue
            if {"lt": a < b, "gt": a > b, "eq": a == b}[op]:
                out.append(r["Code"])
        return out
    for i, op in enumerate(ops):
        assert hits(rows, i) == expect(op, True), op
        assert hits(rows, i + 3) == expect(op, False), op
    by_id = {r["Id"]: r["_rules"] for r in rows}
    assert by_id[ids[0]] == [0, 3] and by_id[ids[1]] == [2, 5] and by_id[ids[2]] == [1, 4]
    assert by_id[ids[3]] == [0]      # boş < 10: boş 0 sayılır; seçenek kapalıyken tutmaz
    assert by_id[ids[4]] == [1]      # 5 > boş
    assert by_id[ids[5]] == [0, 3]   # ondalık


ORDER_LINES = ListDefinition(
    name="pytest_order_lines", table="core.SalesOrderLine",
    fields=("OrderNo", "LineNo", "OpenQty", "RequestedDate", "ConfirmedDate", "UpdatedAt"),
    default_visible=("OrderNo", "LineNo", "RequestedDate", "ConfirmedDate", "UpdatedAt"),
    default_sort=(("OrderNo", "asc"), ("LineNo", "asc")), decimals={"OpenQty": 2})


def local_day(iso: str) -> date:
    return datetime.fromisoformat(iso).astimezone().date()


def test_date_field_ref_compares_days_in_local_time(txn):
    """Tarih–tarih ve zamanlı alan–tarih gün olarak; zamanlı alan (UTC) API sunucusunun yerel gününe çevrilir."""
    conn, session = txn
    ids = conn.execute(text("SELECT TOP 5 Id FROM core.SalesOrderLine WHERE VersionTo IS NULL ORDER BY Id")).scalars().all()
    late_utc = datetime(2026, 10, 28, 23, 30)  # UTC; yerel saat UTC'nin doğusundaysa yerel gün 29 Ekim
    local = late_utc.replace(tzinfo=timezone.utc).astimezone().date()
    data = [(date(2026, 11, 2), date(2026, 11, 5), late_utc), (date(2026, 11, 2), date(2026, 11, 1), late_utc),
            (local, local, late_utc), (date(2026, 11, 2), None, late_utc)]
    for i, (req_d, conf_d, upd) in zip(ids, data):
        conn.execute(text("UPDATE core.SalesOrderLine SET RequestedDate = :r, ConfirmedDate = :c, UpdatedAt = :u "
                          "WHERE Id = :i"), {"r": req_d, "c": conf_d, "u": upd, "i": i})
    req = QueryRequest(fields=list(ORDER_LINES.default_visible), limit=100, rules=[
        rule("ConfirmedDate", op="gt", field_ref="RequestedDate"), rule("ConfirmedDate", op="lt", field_ref="RequestedDate"),
        rule("UpdatedAt", op="eq", field_ref="RequestedDate"), rule("RequestedDate", op="lt", field_ref="ConfirmedDate")])
    rows = {r["Id"]: r for r in list_service.query(session, ORDER_LINES, req)["rows"]}
    assert 0 in rows[ids[0]]["_rules"] and 1 in rows[ids[1]]["_rules"]
    assert 2 in rows[ids[2]]["_rules"]                         # yerel gün eşit
    assert not {0, 1, 3} & set(rows[ids[3]]["_rules"])        # teyit tarihi boş: ne kolon ne karşılaştırılan alan tutar
    for r in rows.values():  # bütün satırlar Python'daki beklentiyle aynı
        rq, cf = r["RequestedDate"], r["ConfirmedDate"]
        assert (0 in r["_rules"]) == bool(rq and cf and cf > rq)
        assert (1 in r["_rules"]) == bool(rq and cf and cf < rq)
        assert (2 in r["_rules"]) == bool(rq and r["UpdatedAt"] and local_day(r["UpdatedAt"]) == date.fromisoformat(rq))
        assert (3 in r["_rules"]) == bool(rq and cf and rq < cf)


def test_date_constant_rules_by_day(api):
    rows = all_rows(api, fields=["Code", "UpdatedAt"])
    day = max(local_day(r["UpdatedAt"]) for r in rows)
    got = all_rows(api, fields=["Code", "UpdatedAt"], rules=[
        rule("UpdatedAt", op="eq", value=day.isoformat()), rule("UpdatedAt", op="lt", value=day.isoformat()),
        rule("UpdatedAt", op="gt", value=day.isoformat())])
    for r in got:
        d = local_day(r["UpdatedAt"])
        assert [i for i in range(3) if i in r["_rules"]] == [i for i, ok in enumerate((d == day, d < day, d > day)) if ok]


# --- görünüm (K-89) ----------------------------------------------------------------------------------------

RULES = [
    {"id": "k1", "field": "SafetyStock", "enabled": True, "condition": {"op": "lt", "field_ref": "MaxStock"},
     "format": {"fill": "#F4CFC4", "color": "#7a2a14", "bold": True, "icon": "arrow_up"}},
    {"id": "k2", "field": "Name", "enabled": False, "condition": {"op": "contains", "value": "mil"},
     "format": {"bold": True}},
    {"id": "k3", "field": "Code", "enabled": True, "condition": {"op": "always"}, "format": {"fill": "#eef0fb"}},
]
VIEW = {"columns": [{"field": "Code"}, {"field": "Name"}, {"field": "SafetyStock"}], "pinned": 0, "sort": [],
        "filters": [], "group_by": []}


def test_rules_saved_with_view(api, users):
    v = create(api, users[0], "pytest Kurallar", {**VIEW, "rules": RULES})
    got = next(x for x in listed(api, users[0])["views"] if x["id"] == v["id"])
    assert got["definition"]["rules"][0]["format"]["fill"] == "#f4cfc4"  # küçük harfe çevrilir
    assert [r["id"] for r in got["definition"]["rules"]] == ["k1", "k2", "k3"] and got["warnings"] == []
    assert got["definition"]["rules"][2]["condition"] == {"op": "always"}


@pytest.mark.parametrize("rules, field", [
    ([{**RULES[0], "format": {}}], "body.definition.rules.0.format"),
    ([{**RULES[0], "format": {"fill": "red"}}], "body.definition.rules.0.format.fill"),
    ([{**RULES[0], "format": {"icon": "smile"}}], "body.definition.rules.0.format.icon"),
    ([RULES[0], {**RULES[1], "id": "k1"}], "definition.rules"),
    ([{**RULES[0], "condition": {"op": "lt", "field_ref": "Name"}}], "definition.rules.0.condition.field_ref"),
])
def test_invalid_rules_in_view_are_422(api, users, rules, field):
    r = api.post(f"/api/materials/views?user_id={users[0]}",
                 json={"name": "pytest Hatalı", "definition": {**VIEW, "rules": rules}})
    assert r.status_code == 422, r.text
    assert any(d["field"].startswith(field) for d in r.json()["error"]["details"]), r.json()


def test_rule_on_removed_field_is_dropped(api, users, db_engine):
    v = create(api, users[0], "pytest Eski kural", {**VIEW, "rules": RULES})
    raw = ('{"v": 1, "columns": [{"field": "Code"}], "rules": ['
           '{"id": "a", "field": "OldField", "condition": {"op": "empty"}, "format": {"bold": true}},'
           '{"id": "b", "field": "SafetyStock", "condition": {"op": "lt", "field_ref": "OldRef"}, "format": {"bold": true}},'
           '{"id": "c", "field": "Name", "condition": {"op": "empty"}, "format": {"bold": true}}]}')
    with db_engine.begin() as conn:
        conn.execute(text("UPDATE adm.TableView SET Definition = :d WHERE Id = :i"), {"d": raw, "i": v["id"]})
    got = next(x for x in listed(api, users[0])["views"] if x["id"] == v["id"])
    assert [r["id"] for r in got["definition"]["rules"]] == ["c"]
    assert got["warnings"] == ["'OldField' alanı artık listede yok; görünümden çıkarıldı.",
                               "'OldRef' alanı artık listede yok; görünümden çıkarıldı."]


# --- Excel (K-92, K-96) ------------------------------------------------------------------------------------

def excel_rule(rid, field, fmt, enabled=True, **condition):
    return {"id": rid, "field": field, "enabled": enabled, "condition": condition, "format": fmt}


def test_excel_cells_follow_rules(api, user):  # noqa: F811
    rules = [
        excel_rule("a", "SafetyStock", {"fill": "#f4cfc4", "icon": "arrow_up"}, op="gt", value=10),
        excel_rule("b", "SafetyStock", {"fill": "#dce0f7", "bold": True}, op="gt", value=0),   # üstteki zemin kazanır
        excel_rule("c", "Name", {"color": "#1e2257"}, op="contains", value="mil"),
        excel_rule("d", "Name", {"bold": True}, enabled=False, op="empty"),                    # pasif
    ]
    body = {"fields": ["Code", "Name", "SafetyStock"], "rules": rules, "group_by": ["MaterialLevel"],
            "sort": [{"field": "Code", "dir": "asc"}]}
    r = api.post(f"/api/materials/export?user_id={user}", json=body)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Veri"]
    screen = {x["Code"]: x for x in all_rows(api, fields=["Code", "Name", "SafetyStock"],
                                              rules=[{k: v for k, v in x.items() if k in ("field", "condition")}
                                                     for x in rules[:3]])}
    seen = 0
    for row in ws.iter_rows(min_row=2):
        code = row[0].value
        if code not in screen:  # grup başlığı: kural uygulanmaz, grup tonunda
            assert row[2].fill.fgColor.rgb not in ("FFF4CFC4", "FFDCE0F7")
            continue
        seen += 1
        on = screen[code]["_rules"]
        stock = row[2]
        if 0 in on:
            assert stock.fill.fgColor.rgb == "FFF4CFC4" and stock.number_format.startswith('"↑"* ')  # ikon başta
            assert stock.font.b and isinstance(stock.value, (int, float))   # kalın alttaki kuraldan; değer sayı
            assert stock.font.color.rgb == "FF1F2430"   # yazı rengi seçilmedi: açık zeminde varsayılan koyu
        elif 1 in on:
            assert stock.fill.fgColor.rgb == "FFDCE0F7" and stock.font.b
        else:
            assert stock.fill.fgColor.rgb in (None, "00000000") and not stock.font.b
        assert (row[1].font.color is not None and row[1].font.color.rgb == "FF1E2257") == (2 in on)
        assert not (row[1].font.b and 2 not in on)  # pasif kural uygulanmadı
    assert seen == len(screen)
    info = dict(wb["Bilgi"].iter_rows(values_only=True))
    assert info["Koşullu biçim"].startswith("Güvenlik stoğu: > 10 → zemin #f4cfc4, ikon yukarı ok;")


def test_excel_field_ref_rule(txn, test_users):
    conn, session = txn
    ids = conn.execute(text("SELECT TOP 3 Id FROM core.Material WHERE VersionTo IS NULL ORDER BY Code")).scalars().all()
    for i, (s, m) in zip(ids, [(5, 10), (20, 10), (None, 10)]):
        conn.execute(text("UPDATE core.Material SET SafetyStock = :s, MaxStock = :m WHERE Id = :i"), {"s": s, "m": m, "i": i})
    req = ExportRequest(fields=["Code", "SafetyStock", "MaxStock"], sort=[{"field": "Code", "dir": "asc"}], rules=[
        excel_rule("x", "SafetyStock", {"fill": "#f4cfc4"}, op="lt", field_ref="MaxStock")])
    out = export_service.start(session, MATERIALS, test_users[0], req)
    ws = load_workbook(io.BytesIO(out["file"]))["Veri"]
    rows = list(ws.iter_rows(min_row=2, max_row=4))
    assert [r[1].fill.fgColor.rgb == "FFF4CFC4" for r in rows] == [True, False, True]  # boş < 10 (boşlar 0)
    assert rows[0][1].value == 5


def test_excel_auto_text_color_and_icon_side(api, user):  # noqa: F811
    """Koyu zeminde yazı rengi seçilmemişse beyaz (ekranla aynı hesap); ikon metinde sonda, sayıda başta (K-96)."""
    rules = [excel_rule("n", "Name", {"fill": "#002060", "icon": "flag", "icon_color": "#c00000"}, op="contains", value="a"),
             excel_rule("s", "SafetyStock", {"icon": "warning"}, op="empty")]
    r = api.post(f"/api/materials/export?user_id={user}",
                 json={"fields": ["Code", "Name", "SafetyStock"], "rules": rules, "sort": [{"field": "Code", "dir": "asc"}]})
    ws = load_workbook(io.BytesIO(r.content))["Veri"]
    names = [row[1] for row in ws.iter_rows(min_row=2) if row[1].fill.fgColor.rgb == "FF002060"]
    assert names and all(c.font.color.rgb == "FFFFFFFF" for c in names)        # okunur: beyaz
    assert all(c.number_format == '@* "⚑"' for c in names)                     # metin: ikon sonda
    stocks = [row[2] for row in ws.iter_rows(min_row=2) if row[2].value is None and row[2].number_format != "General"]
    assert stocks and all(c.number_format.startswith('"▲"* ') for c in stocks)  # sayı: ikon başta (uyarı ▲)


def test_excel_always_rule_formats_every_cell(api, user):  # noqa: F811
    """"Her zaman" kolon biçimi Excel'de kolonun bütün veri hücrelerine (boşlar dahil) uygulanır; grup satırına değil."""
    rules = [excel_rule("k", "SafetyStock", {"fill": "#eef0fb", "bold": True}, op="always")]
    r = api.post(f"/api/materials/export?user_id={user}",
                 json={"fields": ["Code", "SafetyStock"], "rules": rules, "group_by": ["MaterialLevel"]})
    wb = load_workbook(io.BytesIO(r.content))
    data = [row for row in wb["Veri"].iter_rows(min_row=2) if row[0].value and not str(row[0].value).endswith(" satır")]
    assert len(data) == 200
    assert all(row[1].fill.fgColor.rgb == "FFEEF0FB" and row[1].font.b for row in data)
    assert any(row[1].value is None for row in data)  # boş hücre de boyalı
    assert dict(wb["Bilgi"].iter_rows(values_only=True))["Koşullu biçim"] == "Güvenlik stoğu: kolon biçimi → zemin #eef0fb, kalın"
