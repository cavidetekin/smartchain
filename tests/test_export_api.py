"""Excel'e aktarım (K-91, K-93; fonksiyon 86): sentetik v0 ve yerel SQL Server ile.

Anlık aktarım ekrandaki sorguyla aynı satırları aynı sırada verir; eşik üstü "Dışa aktarım" kanalında iş olur, motor
servisi dosyayı üretir; her aktarım AuditLog'a yazılır. Test kullanıcısının iz satırları sonra silinir.
"""
import io
import json
import os
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import bindparam, text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.orm import Session

from app.api.db import get_session
from app.api.main import create_app
from app.api.repositories import export as export_repo
from app.api.services import export as export_service
from app.engine.export_job import run_export_job
from app.engine.service import EngineService
from app.jobs import queue
from app.jobs.constants import (CHANNEL_EXPORT, CHANNEL_INTERACTIVE, STATUS_DONE, STATUS_FAILED, STATUS_QUEUED, TYPE_EXPORT,
                                TYPE_WORKSPACE)
from app.jobs.settings import JobSettings

from test_lists_api import all_rows, api  # noqa: F401  (fikstür ve yardımcı)

pytestmark = pytest.mark.db

FIELDS = ["Code", "Name", "MaterialGroup", "SafetyStock", "ProductionLeadTimeDays", "UpdatedAt", "IsPurchasable"]
SORT = [{"field": "MaterialGroup", "dir": "desc"}, {"field": "Code", "dir": "asc"}]
FILTERS = [{"field": "Name", "op": "contains", "value": "mil"}]


@pytest.fixture
def user(db_engine, test_users, clean_jobs):
    """Deneme kullanıcısı; iz satırları, dışa aktarım ayarları ve iş tablosu her testten sonra temizlenir."""
    def clean():
        with db_engine.begin() as c:
            c.execute(text("DELETE FROM adm.AuditLog WHERE UserId IN :ids").bindparams(bindparam("ids", expanding=True)),
                      {"ids": test_users})
            c.execute(text("DELETE FROM adm.SystemSetting WHERE SettingKey LIKE 'Export%'"))
    clean()
    yield test_users[0]
    clean()


def setting(db_engine, key, value):
    with db_engine.begin() as c:
        c.execute(text("DELETE FROM adm.SystemSetting WHERE SettingKey = :k"), {"k": key})
        c.execute(text("INSERT INTO adm.SystemSetting (SettingKey, Value) VALUES (:k, :v)"), {"k": key, "v": str(value)})


def export(api, user, **body):
    return api.post(f"/api/materials/export?user_id={user}", json={"fields": FIELDS, **body})


def sheet(content: bytes):
    wb = load_workbook(io.BytesIO(content))
    return wb, [list(r) for r in wb["Veri"].iter_rows(values_only=True)]


def audits(db_engine, user):
    with db_engine.connect() as c:
        return c.execute(text("SELECT ObjectType, ObjectKey, NewValue FROM adm.AuditLog WHERE UserId = :u "
                              "AND Action = N'Dışa aktarım' ORDER BY Id"), {"u": user}).all()


def test_sync_export_matches_screen_query(api, user, db_engine):
    r = export(api, user, sort=SORT, filters=FILTERS, widths={"Code": 8})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == export_service.XLSX
    disposition = r.headers["content-disposition"]
    assert disposition.startswith("attachment; filename*=UTF-8''smartchain_malzeme_veri-v") and disposition.endswith(".xlsx")
    wb, rows = sheet(r.content)
    screen = all_rows(api, fields=FIELDS, sort=SORT, filters=FILTERS)
    labels = {c["field"]: c["label"] for c in api.get("/api/materials/columns").json()["columns"]}
    assert rows[0] == [labels[f] for f in FIELDS]
    assert [r[0] for r in rows[1:]] == [s["Code"] for s in screen] and len(screen) > 0
    for row, s in zip(rows[1:], screen):  # sayılar sayı olarak, ekrandaki değerle aynı
        assert row[3] == (None if s["SafetyStock"] is None else pytest.approx(float(s["SafetyStock"])))
        assert row[4] == s["ProductionLeadTimeDays"]
    assert all(r[6] in ("Evet", "Hayır") for r in rows[1:])
    info = dict(wb["Bilgi"].iter_rows(values_only=True))
    assert info["Filtreler"] == "Malzeme adı: içerir 'mil'"
    assert info["Satırlar"] == f"Filtrelenmiş · {len(screen)}"
    assert info["Sıralama"].startswith("Malzeme grubu (azalan)")
    [(obj, key, value)] = audits(db_engine, user)
    details = json.loads(value)
    assert (obj, details["rows"], details["background"]) == ("materials", len(screen), False)
    assert disposition.endswith(key.replace(" ", "%20"))


def test_columns_carry_export_info(api, user, db_engine):
    setting(db_engine, "ExportSyncMaxRows", 2500)
    assert api.get("/api/materials/columns").json()["export"] == {
        "title": "Malzeme listesi", "file_slug": "malzeme", "sync_max_rows": 2500, "max_rows": 500000, "file_hours": 24}


def test_scope_all_rows_all_columns(api, user):
    # Eski istemcinin gönderdiği kaldırılmış seçenekler yok sayılır: Bilgi sayfası her zaman (K-91).
    r = export(api, user, filters=FILTERS, scope={"rows": "all", "columns": "all", "info_sheet": False, "keep_groups": False})
    wb, rows = sheet(r.content)
    assert wb.sheetnames == ["Veri", "Bilgi"]
    assert len(rows) - 1 == 200  # sentetik v0: bütün malzemeler
    assert len(rows[0]) > len(FIELDS) and rows[0][: len(FIELDS)][0] == "Malzeme kodu"


def test_grouped_export_follows_group_tree(api, user):
    body = {"sort": [{"field": "Code", "dir": "asc"}], "group_by": ["MaterialLevel", "MaterialGroup"],
            "aggregates": {"SafetyStock": {"fn": "sum"}}}
    r = export(api, user, fields=["Name", "Code", "SafetyStock"], **body)
    assert r.status_code == 200, r.text
    _, rows = sheet(r.content)
    tree = api.post("/api/materials/groups", json={k: v for k, v in body.items()}).json()

    expected = []
    def walk(nodes, depth):
        for n in nodes:
            expected.append((depth, n["count"], n["aggregates"].get("SafetyStock")))
            walk(n.get("children", []), depth + 1)
    walk(tree["groups"], 0)

    groups, leaf_counts, current = [], [], None
    for row in rows[1:]:
        if row[1] is None and isinstance(row[0], str) and row[0].endswith(" satır"):
            groups.append(row)
            current = [0]
            leaf_counts.append(current)
        else:
            current[0] += 1
    assert len(groups) == len(expected)
    assert [int(g[0].rsplit(" · ", 1)[1].split()[0].replace(".", "")) for g in groups] == [e[1] for e in expected]
    leaf_expected = [e[1] for i, e in enumerate(expected) if e[0] == 1]
    assert [c[0] for c, e in zip(leaf_counts, expected) if e[0] == 1] == leaf_expected
    assert sum(leaf_expected) == 200


def test_trace_failure_gives_no_file(api, user, monkeypatch):
    def broken(*a, **k):
        raise DBAPIError("INSERT", {}, Exception("AuditLog yazılamadı"))
    monkeypatch.setattr(export_repo, "write_audit", broken)
    r = export(api, user)
    assert r.status_code == 503 and r.json()["error"]["code"] == "export_trace_failed"


def test_validation_and_limits(api, user, db_engine):
    r = export(api, user, fields=["Secret"])
    assert r.status_code == 422 and r.json()["error"]["code"] == "validation"
    assert api.post(f"/api/materials/export?user_id={user}", json={"fields": []}).status_code == 422
    assert export(api, 987654).status_code == 404
    setting(db_engine, "ExportMaxRows", 10)
    r = export(api, user)
    assert r.status_code == 422 and r.json()["error"]["code"] == "export_too_large"
    assert "200 satır" in r.json()["error"]["message"]
    assert audits(db_engine, user) == []


def test_database_unavailable_is_503():
    from fastapi.testclient import TestClient
    app = create_app()
    app.dependency_overrides[get_session] = lambda: None
    c = TestClient(app, raise_server_exceptions=False)
    assert c.post("/api/materials/export?user_id=1", json={"fields": ["Code"]}).status_code == 503
    assert c.get("/api/exports/1/file?user_id=1").status_code == 503


def enqueue_big(api, user, db_engine):
    setting(db_engine, "ExportSyncMaxRows", 10)
    r = export(api, user, sort=SORT, filters=[{"field": "Name", "op": "not_contains", "value": "zzz"}])
    assert r.status_code == 202, r.status_code
    return r.json()


def test_above_threshold_becomes_export_job(api, user, db_engine, db_session):
    out = enqueue_big(api, user, db_engine)
    assert out["rows"] > 10 and out["file_name"].startswith("smartchain_malzeme_")
    job = queue.get_job(db_session, out["job_id"], JobSettings())
    assert (job["JobType"], job["Channel"], job["Status"], job["RequestedBy"]) == (TYPE_EXPORT, CHANNEL_EXPORT,
                                                                                      STATUS_QUEUED, user)
    assert json.loads(job["Payload"])["request"]["filters"][0]["value"] == "zzz"
    r = api.get(f"/api/exports/{out['job_id']}/file?user_id={user}")
    assert r.status_code == 409 and r.json()["error"]["code"] == "export_not_ready"
    assert audits(db_engine, user) == []  # iz dosya hazır olunca


def test_export_job_writes_file_audit_and_download(api, user, db_engine, db_session, test_users):
    out = enqueue_big(api, user, db_engine)
    job = queue.claim_next(db_session, CHANNEL_EXPORT, JobSettings())
    assert job["Id"] == out["job_id"]
    run_export_job(job["Id"], job["Attempt"])
    db_session.expire_all()
    done = queue.get_job(db_session, job["Id"], JobSettings())
    assert done["Status"] == STATUS_DONE and done["Progress"].startswith("Hazır · ")
    path = export_service.export_path(job["Id"])
    assert path.exists() and not path.with_suffix(".tmp").exists()
    [(_, key, value)] = audits(db_engine, user)
    assert key == out["file_name"] and json.loads(value)["background"] is True

    r = api.get(f"/api/exports/{job['Id']}/file?user_id={user}")
    assert r.status_code == 200 and r.headers["content-type"] == export_service.XLSX
    assert out["file_name"] in r.headers["content-disposition"] or "filename*=" in r.headers["content-disposition"]
    _, rows = sheet(r.content)
    assert len(rows) - 1 == out["rows"]

    r = api.get(f"/api/exports/{job['Id']}/file?user_id={test_users[1]}")
    assert r.status_code == 403 and r.json()["error"]["code"] == "export_not_owner"
    with db_engine.begin() as c:
        c.execute(text("UPDATE adm.Job SET EndedAt = DATEADD(hour, -25, EndedAt) WHERE Id = :id"), {"id": job["Id"]})
    r = api.get(f"/api/exports/{job['Id']}/file?user_id={user}")
    assert r.status_code == 410 and r.json()["error"]["code"] == "export_expired"
    assert api.get(f"/api/exports/987654/file?user_id={user}").status_code == 404


def test_export_job_failure_marks_job_failed(api, user, db_engine, db_session, monkeypatch):
    out = enqueue_big(api, user, db_engine)
    job = queue.claim_next(db_session, CHANNEL_EXPORT, JobSettings())
    monkeypatch.setattr(export_repo, "write_audit", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("iz yok")))
    run_export_job(job["Id"], job["Attempt"])
    db_session.expire_all()
    failed = queue.get_job(db_session, out["job_id"], JobSettings())
    assert failed["Status"] == STATUS_FAILED and str(out["job_id"]) in failed["ErrorText"]
    assert not export_service.export_path(job["Id"]).exists()  # iz yoksa dosya yok


def test_export_channel_has_own_limit(user, db_session):
    settings = JobSettings(limit_export=2, limit_interactive=1)
    for _ in range(3):
        queue.enqueue(db_session, TYPE_EXPORT, CHANNEL_EXPORT, user, {})
    queue.enqueue(db_session, TYPE_WORKSPACE, CHANNEL_INTERACTIVE, user, {})
    claimed = [queue.claim_next(db_session, CHANNEL_EXPORT, settings) for _ in range(3)]
    assert [j is not None for j in claimed] == [True, True, False]
    assert queue.claim_next(db_session, CHANNEL_INTERACTIVE, settings) is not None  # aktarım yeniden hesabı bekletmez
    c = queue.counters(db_session, settings)
    assert c[CHANNEL_EXPORT] == {"running": 2, "queued": 1, "limit": 2}


def test_engine_service_dispatches_export_job_by_type(api, user, db_engine, db_session):
    out = enqueue_big(api, user, db_engine)
    svc = EngineService(db_engine, job_settings=JobSettings(), channels=(CHANNEL_EXPORT,))
    svc.start()
    try:
        end = time.monotonic() + 60
        while time.monotonic() < end:
            svc.poll_once()
            db_session.expire_all()
            job = queue.get_job(db_session, out["job_id"], JobSettings())
            if job["Status"] == STATUS_DONE and not svc.running:
                break
            time.sleep(0.2)
        assert job["Status"] == STATUS_DONE, job
    finally:
        svc.shutdown()
    assert export_service.export_path(out["job_id"]).exists()


def test_cleanup_removes_expired_files(db_engine, user):
    folder = Path(export_service.export_path(1).parent)
    folder.mkdir(parents=True, exist_ok=True)
    old, new = folder / "990001.xlsx", folder / "990002.xlsx"
    for f in (old, new):
        f.write_bytes(b"x")
    past = time.time() - 25 * 3600
    os.utime(old, (past, past))
    with Session(db_engine) as s:
        removed = EngineService(db_engine, job_settings=JobSettings()).cleanup_exports(s)
    assert "990001.xlsx" in removed and not old.exists() and new.exists()
    new.unlink()
