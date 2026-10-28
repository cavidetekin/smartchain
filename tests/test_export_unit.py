"""Excel yazıcısı ve biçim tanımı (K-91, K-92): veritabanısız birim testleri."""
import io
import json
from datetime import date, datetime, timezone
from decimal import Decimal

from openpyxl import load_workbook

from app.api.export import describe, formats
from app.api.export.formats import CellFormat
from app.api.export.writer import XlsxWriter, local_datetime, width_chars

COLUMNS = [
    {"field": "Code", "label": "Malzeme kodu", "type": "text"},
    {"field": "Qty", "label": "Güvenlik stoğu", "type": "number", "decimals_from_unit": True},
    {"field": "Days", "label": "Üretim süresi", "type": "integer"},
    {"field": "Pct", "label": "Fire", "type": "number", "decimals": 1},
    {"field": "Flag", "label": "Satın alınabilir", "type": "bool"},
    {"field": "Day", "label": "Termin", "type": "date"},
    {"field": "At", "label": "Son değişiklik", "type": "datetime"},
]
ROW = {"Code": "M-1", "Qty": Decimal("1234.5"), "Days": 3, "Pct": 2.25, "Flag": True,
       "Day": date(2026, 10, 28), "At": datetime(2026, 10, 28, 12, 30)}


class Book:
    """Yazıcıyı bellekteki dosyaya açar; read() kaydedip openpyxl ile okur."""

    def __init__(self, columns=COLUMNS, widths=None, **kw):
        self.buf = io.BytesIO()
        self.w = XlsxWriter(self.buf, columns, widths, **kw)

    def __getattr__(self, name):  # header, data_row, group_row, info_sheet
        return getattr(self.w, name)

    def read(self):
        self.w.save()
        self.buf.seek(0)
        return load_workbook(self.buf)


def test_values_are_typed_not_text():
    w = Book(widths={"Code": 8})
    w.header()
    w.data_row(ROW, unit_decimals=3)
    ws = w.read()["Veri"]
    cells = list(ws.iter_rows(min_row=2, max_row=2))[0]
    values = [c.value for c in cells]
    assert values[:5] == ["M-1", 1234.5, 3, 2.25, "Evet"]
    assert isinstance(values[1], float) and isinstance(values[2], int)
    assert values[5] == datetime(2026, 10, 28)
    assert values[6] == local_datetime(ROW["At"])
    assert [c.number_format for c in cells[1:]] == ["#,##0.000", "#,##0", "#,##0.0", "General", "dd.mm.yyyy",
                                                     "dd.mm.yyyy hh:mm"]


def test_empty_values_are_empty_cells():
    w = Book()
    w.header()
    w.data_row({"Code": "", "Qty": None, "Flag": False})
    values = [c.value for c in list(w.read()["Veri"].iter_rows(min_row=2))[0]]
    assert values == [None, None, None, None, "Hayır", None, None]


def test_header_frozen_filter_widths_and_colors():
    w = Book(widths={"Code": 7, "Qty": 14})
    w.header()
    w.data_row(ROW)
    ws = w.read()["Veri"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:G2"
    # Excel genişliği karakter dolgusuyla saklar (13 karakter → 13,71); Excel'de görünen değer width_chars'tır.
    shown = lambda c: round(ws.column_dimensions[c].width - 0.71, 1)  # noqa: E731
    assert shown("A") == width_chars(7) == 13.0
    assert shown("C") == width_chars(None)
    head = ws["A1"]
    assert head.value == "Malzeme kodu" and head.font.b
    assert head.fill.fgColor.rgb == "FF" + formats.color("table-head-bg")[1:].upper()


def test_group_rows_carry_aggregates_with_screen_tones():
    w = Book()
    w.header()
    w.group_row("Hammadde", 0, 80, {"Qty": 120.0, "Days": 2.5}, {"Qty": "sum", "Days": "avg"}, {"Qty": 0}, "80 satır")
    w.group_row("Çubuk", 1, 20, {"Qty": None}, {"Qty": "sum"}, {}, "20 satır")
    w.data_row(ROW, indent=2)
    ws = w.read()["Veri"]
    g1, g2, data = ws[2], ws[3], ws[4]
    assert [g1[0].value, g1[1].value, g1[2].value] == ["Hammadde · 80 satır", 120, 2.5]
    assert (g1[1].number_format, g1[2].number_format) == ("#,##0", "#,##0.0")  # ortalama bir ondalık fazla
    assert g1[0].fill.fgColor.rgb == "FF" + formats.color("table-group-1-bg")[1:].upper() and g1[0].font.b
    assert g2[0].fill.fgColor.rgb == "FF" + formats.color("table-group-2-bg")[1:].upper() and not g2[0].font.b
    assert (g1[0].alignment.indent, g2[0].alignment.indent, data[0].alignment.indent) == (0, 1, 2)
    assert g2[1].value is None


def test_cell_format_hook_applies_to_data_cells():
    """1.6 (U-24) koşullu biçimleri bu kancadan uygular; şimdiden zemin, yazı rengi ve kalın hücreye işlenir."""
    red = CellFormat(fill="#f4cfc4", color="#7a2a14", bold=True)
    w = Book(cell_format=lambda col, row, v: red if col["field"] == "Qty" and v > 1000 else None)
    w.header()
    w.data_row(ROW)
    cell = w.read()["Veri"]["B2"]
    assert (cell.fill.fgColor.rgb, cell.font.color.rgb, cell.font.b) == ("FFF4CFC4", "FF7A2A14", True)
    assert cell.number_format == "#,##0.00"  # birim bilinmiyorsa 2 ondalık


def test_info_sheet():
    w = Book()
    w.header()
    w.info_sheet([("Liste", "Malzeme listesi"), ("Filtreler", "yok")])
    wb = w.read()
    assert wb.sheetnames == ["Veri", "Bilgi"]
    assert [tuple(r) for r in wb["Bilgi"].iter_rows(values_only=True)] == [
        ("Alan", "Değer"), ("Liste", "Malzeme listesi"), ("Filtreler", "yok")]


def test_cell_format_merge():
    base = CellFormat(fill="#ffffff", bold=True)
    assert base.merged(CellFormat(color="#000000")) == CellFormat(fill="#ffffff", color="#000000", bold=True)
    assert base.merged(None) is base


def test_tokens_cover_formats():
    """Biçim tanımının okuduğu renkler design-tokens.json'da var (tokens.css'ten üretilir; eşitliği vitest denetler)."""
    data = json.loads(formats.TOKENS.read_text(encoding="utf-8"))["colors"]
    for name in ("table-head-bg", "table-head-text", "table-group-1-bg", "table-group-2-bg", "text"):
        assert data[name].startswith("#") and len(data[name]) == 7
    for level in ("low", "normal", "high", "risk", "over"):
        assert formats.occupancy_format(level).fill == data[f"occ-{level}-bg"]


def test_aggregate_formats():
    col = {"type": "number", "decimals_from_unit": True}
    assert formats.aggregate_number_format(col, "sum", 3) == "#,##0.000"
    assert formats.aggregate_number_format(col, "wavg", 0) == "#,##0.0"
    assert formats.aggregate_number_format(col, "count", 3) == "#,##0"


def test_describe_texts_are_turkish():
    assert describe.tr_number(1234567.5, 2) == "1.234.567,50"
    assert describe.tr_number(20, 0) == "20"
    assert describe.describe_filter({"op": "contains", "value": "mil"}, "Malzeme adı") == "Malzeme adı: içerir 'mil'"
    assert describe.describe_filter({"op": "in", "values": ["Çubuk", True], "empty": True}, "Grup") == "Grup: Çubuk, Evet, boş"
    assert describe.describe_filter({"op": "between", "min": 10, "max": 1000.5}, "Stok", 1) == "Stok: 10,0–1.000,5"
    assert describe.describe_filter({"op": "between", "from": "2026-10-01"}, "Termin") == "Termin: 01.10.2026 ve sonrası"
    assert describe.describe_filter({"op": "relative", "value": "this_week"}, "Değişiklik") == "Değişiklik: bu hafta"
    assert describe.describe_sort([{"field": "Code", "dir": "desc"}], {"Code": "Kod"}) == "Kod (azalan)"
    assert describe.describe_aggregates({"Q": {"fn": "wavg", "weight": "W"}}, {"Q": "Miktar", "W": "Ağırlık"}) == \
        "Miktar: ağırlıklı ortalama (ağırlık: Ağırlık)"


def test_local_datetime_from_utc():
    utc = datetime(2026, 1, 1, 10, 0)
    assert local_datetime(utc) == utc.replace(tzinfo=timezone.utc).astimezone().replace(tzinfo=None)


def test_text_is_never_formula_or_link():
    """Veriden gelen metin formül ya da bağlantı olmaz (formül enjeksiyonu yok; K-91)."""
    w = Book()
    w.header()
    w.data_row({"Code": "=HYPERLINK(\"x\",\"tıkla\")", "Flag": True})
    w.data_row({"Code": "www.ornek.local/a"})
    ws = w.read()["Veri"]
    assert ws["A2"].value == "=HYPERLINK(\"x\",\"tıkla\")" and ws["A2"].data_type == "s"
    assert ws["A3"].value == "www.ornek.local/a" and ws["A3"].hyperlink is None
