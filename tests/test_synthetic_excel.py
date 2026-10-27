"""Kirli hâl: entegrasyon biçiminde Excel ve beklenen hatalar (docs/11, K-78). Veritabanı gerekmez."""
import json
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from app.data.synthetic import dataset as d
from app.data.synthetic.dataset import generate
from app.data.synthetic.excel import FIRST_DATA_ROW, stg_columns, stg_entities, to_staging, write_excel

DAY = date(2026, 10, 27)


@pytest.fixture(scope="module")
def written(tmp_path_factory):
    out = tmp_path_factory.mktemp("kirli")
    ds = generate(DAY)
    manifest = write_excel(ds, out)
    return ds, out, manifest


def read(out, entity):
    ws = load_workbook(out / f"{entity}.xlsx", read_only=True)[entity]
    rows = list(ws.iter_rows(values_only=True))
    header, data = rows[0], rows[FIRST_DATA_ROW - 1:]
    return list(header), rows[1], [dict(zip(header, r)) for r in data]


def norm(v):
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, datetime) and v.time() == datetime.min.time():
        return v.date()
    return v


def test_one_file_per_integration_entity_with_stg_headers(written):
    ds, out, manifest = written
    files = sorted(p.name for p in out.glob("*.xlsx"))
    assert files == sorted(f"{e}.xlsx" for e in stg_entities()) and len(files) == 25
    for entity in stg_entities():
        header, descriptions, data = read(out, entity)
        assert header == [c.name for c in stg_columns(entity)]
        assert all(desc for desc in descriptions)
        assert len(data) == manifest["files"][f"{entity}.xlsx"]
    assert (out / "beklenen-hatalar.json").exists()


def test_expected_errors_are_where_listed(written):
    ds, out, manifest = written
    saved = json.loads((out / "beklenen-hatalar.json").read_text(encoding="utf-8"))
    assert {e["code"] for e in saved["errors"]} == {"H1", "H2", "H3", "H4"}
    for e in saved["errors"]:
        header, _, data = read(out, e["file"].removesuffix(".xlsx"))
        for row in e["rows"]:
            assert norm(data[row - FIRST_DATA_ROW][e["field"]]) == e["value"], e
        assert e["severity"] in ("hata", "uyarı") and e["message"] and e["step"]
    h4 = next(e for e in saved["errors"] if e["code"] == "H4")
    _, _, mats = read(out, "Material")
    a, b = (mats[r - FIRST_DATA_ROW] for r in h4["rows"])
    assert a["Code"] == b["Code"] == d.ISOLATED_MATERIAL and a["Name"] != b["Name"]


def test_isolated_and_missing_route(written):
    ds, out, _ = written
    _, _, routing = read(out, "Routing")
    assert not [r for r in routing if r["RouteCode"] == f"R-{d.NO_ROUTE_PART}"]
    _, _, pv = read(out, "ProductionVersion")
    assert [r["RouteCode"] for r in pv if r["MaterialCode"] == d.NO_ROUTE_PART] == [f"R-{d.NO_ROUTE_PART}"]
    for entity in stg_entities():  # H4'teki malzeme başka dosyada geçmez: zincirleme hata yok
        if entity != "Material":
            _, _, data = read(out, entity)
            assert d.ISOLATED_MATERIAL not in {v for r in data for v in r.values()}, entity


def test_other_rows_equal_clean_state(written):
    """Hatalı satırlar dışındaki her satır temiz hâlin stg biçimiyle aynı: ETL sonrası = temiz − beklenen fark."""
    ds, out, manifest = written
    clean = to_staging(ds)
    error_rows = {}
    for e in manifest["errors"]:
        if e["severity"] == "hata":  # uyarılı satır (H1) yüklenir, temiz hâlde de vardır
            error_rows.setdefault(e["file"].removesuffix(".xlsx"), set()).update(e["rows"])
    for entity in stg_entities():
        _, _, data = read(out, entity)
        kept = [{k: norm(v) for k, v in r.items()} for i, r in enumerate(data, FIRST_DATA_ROW)
                if i not in error_rows.get(entity, set())]
        expected = [{k: norm(v) for k, v in r.items()} for r in clean[entity]]
        if entity == "Routing":
            expected = [r for r in expected if r["RouteCode"] != f"R-{d.NO_ROUTE_PART}"]
        if entity in ("PurchaseOrderLine", "BOM", "Material"):  # değiştirilen satırlar temizden düşer
            dropped = {("PurchaseOrderLine", "PO-2060"), ("BOM", "ALT-012|SAT-010"), ("Material", d.ISOLATED_MATERIAL)}
            keyf = {"PurchaseOrderLine": lambda r: r["OrderNo"],
                    "BOM": lambda r: f"{r['MaterialCode']}|{r['ComponentMaterialCode']}",
                    "Material": lambda r: r["Code"]}[entity]
            expected = [r for r in expected if (entity, keyf(r)) not in dropped]
        assert kept == expected, entity


def test_codes_resolve_outside_error_rows(written):
    ds, out, manifest = written
    _, _, mats = read(out, "Material")
    material_codes = {r["Code"] for r in mats}
    bad = {e["value"] for e in manifest["errors"] if e["code"] == "H2"}
    for entity in stg_entities():
        _, _, data = read(out, entity)
        for r in data:
            for k in ("MaterialCode", "ComponentMaterialCode"):
                if r.get(k) is not None and r[k] not in bad:
                    assert r[k] in material_codes, (entity, k, r[k])
