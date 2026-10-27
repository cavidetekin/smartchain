"""Kirli hâl: entegrasyon biçiminde Excel ve beklenen hatalar (docs/11, K-78, K-79).

Temiz hâl stg.* alanlarına çevrilir (Id → kod), kasıtlı hatalar eklenir; varlık başına bir dosya (tek dosya tek
varlık). Sayfa düzeni (geçici, 3.1'deki şablonla kesinleşir): 1. satır alan adı, 2. satır açıklama, veri 3. satırdan.
Hatalar başka satırları zincirleme etkilemeyecek yerlerdedir: ETL sonrası ana veri = temiz hâl − beklenen fark.
"""
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from db import models
from . import dataset as d
from .dataset import Dataset

FIRST_DATA_ROW = 3
EMPTY_ENTITIES = {"Forecast"}  # docs/11: tahmin üretilmez

STG_EXCLUDED = {"StagingRowId", "ImportBatchId", "RowStatus", "ErrorText"}


def stg_columns(entity: str) -> list:
    t = models.metadata.tables[f"stg.{entity}"]
    return [c for c in t.columns if c.name not in STG_EXCLUDED]


def stg_entities() -> list[str]:
    return [t.name for t in models.metadata.sorted_tables if t.schema == "stg" and t.name not in EMPTY_ENTITIES]


def _resolvers(ds: Dataset) -> dict:
    """stg kod alanı → temiz satırdan kodu üreten işlev."""
    c = ds.code
    pv = lambda r: ds.by_id("ProductionVersion", r["ProductionVersionId"])  # noqa: E731

    def wo(field):
        return lambda r: None if r.get(field) is None else ds.by_id("WorkOrder", r[field])["OrderNo"]

    def labor_ref(r):
        if r.get("LaborRefId") is None:
            return None
        table = "CompetenceGroup" if r.get("LaborType") == "Yetkinlik grubu" else "ResourceGroup"
        return c(table, r["LaborRefId"])

    def so(field):
        return lambda r: None if r.get("SalesOrderLineId") is None else ds.by_id("SalesOrderLine", r["SalesOrderLineId"])[field]

    return {
        "LocationCode": lambda r: d.LOCATION_CODE,
        "MaterialCode": lambda r: c("Material", r["MaterialId"] if "MaterialId" in r else pv(r)["MaterialId"]),
        "ComponentMaterialCode": lambda r: c("Material", r["ComponentMaterialId"]),
        "ProductionVersionCode": lambda r: pv(r)["Code"],
        "SupplierCode": lambda r: c("Supplier", r.get("SupplierId")),
        "CustomerCode": lambda r: c("Customer", r.get("CustomerId")),
        "ProjectCode": lambda r: c("Project", r.get("ProjectId")),
        "ProjectGroupCode": lambda r: c("ProjectGroup", r.get("ProjectGroupId")),
        "ResourceCode": lambda r: c("Resource", r.get("ResourceId")),
        "ResourceGroupCode": lambda r: c("ResourceGroup", r.get("ResourceGroupId")),
        "WorkerCode": lambda r: c("Worker", r["WorkerId"]),
        "CompetenceGroupCode": lambda r: c("CompetenceGroup", r["CompetenceGroupId"]),
        "WorkOrderNo": wo("WorkOrderId"),
        "ParentWorkOrderNo": wo("ParentWorkOrderId"),
        "SalesOrderNo": so("OrderNo"),
        "SalesOrderLineNo": so("LineNo"),
        "LaborRefCode": labor_ref,
    }


def to_staging(ds: Dataset) -> dict[str, list[dict]]:
    """Temiz hâlin stg biçimi (kodlarla), varlık başına satırlar."""
    resolve = _resolvers(ds)
    out = {}
    for entity in stg_entities():
        rows = []
        for r in ds.rows(entity):
            row = {}
            for col in stg_columns(entity):
                if col.name in resolve:
                    row[col.name] = resolve[col.name](r)
                else:
                    row[col.name] = r.get(col.name)
            rows.append(row)
        out[entity] = rows
    return out


def _find(rows: list[dict], **match) -> int:
    hits = [i for i, r in enumerate(rows) if all(r.get(k) == v for k, v in match.items())]
    if len(hits) != 1:
        raise KeyError(f"{match}: {len(hits)} satır")
    return hits[0]


def make_dirty(ds: Dataset) -> tuple[dict[str, list[dict]], list[dict]]:
    """Kasıtlı hatalı stg satırları ve beklenen sonuçlar (docs/11 · Kirli hâl)."""
    stg = to_staging(ds)
    errors = []

    # H1 · eksik rota: sürüm rotayı gösteriyor, rota dosyada yok.
    route = f"R-{d.NO_ROUTE_PART}"
    removed = [r for r in stg["Routing"] if r["RouteCode"] == route]
    stg["Routing"] = [r for r in stg["Routing"] if r["RouteCode"] != route]

    # H2 · tanımsız malzeme kodu (fazladan satırlar).
    so = dict(stg["SalesOrderLine"][_find(stg["SalesOrderLine"], OrderNo="SO-1001")])
    so.update(OrderNo="SO-1099", MaterialCode="MAM-099")
    stg["SalesOrderLine"].append(so)
    bom_extra = dict(stg["BOM"][_find(stg["BOM"], MaterialCode="ALT-005", ComponentMaterialCode="PRC-013")])
    bom_extra.update(ComponentMaterialCode="SAT-999", Quantity=4.0)
    stg["BOM"].append(bom_extra)

    # H3 · negatif miktar (var olan satırlarda).
    po_i = _find(stg["PurchaseOrderLine"], OrderNo="PO-2060", LineNo=1)
    stg["PurchaseOrderLine"][po_i]["OpenQty"] = -50.0
    bom_i = _find(stg["BOM"], MaterialCode="ALT-012", ComponentMaterialCode="SAT-010")
    stg["BOM"][bom_i]["Quantity"] = -2.0

    # H4 · tekrar eden anahtar: aynı kod iki satır, adları farklı.
    mat_i = _find(stg["Material"], Code=d.ISOLATED_MATERIAL)
    dup = dict(stg["Material"][mat_i], Name="Etiket seti · müşteri temin (eski kayıt)")
    stg["Material"].insert(mat_i + 1, dup)

    row = lambda entity, **m: _find(stg[entity], **m) + FIRST_DATA_ROW  # noqa: E731
    pv_row = row("ProductionVersion", MaterialCode=d.NO_ROUTE_PART)
    errors += [
        dict(code="H1", type="Eksik rota", file="ProductionVersion.xlsx", rows=[pv_row],
             key=f"{d.NO_ROUTE_PART}|V1", field="RouteCode", value=route,
             step="Fonksiyon 20 Veri kalitesi (geliştirme 6.5); ön hesapta da uyarı (fonksiyon 29, 10.1)",
             severity="uyarı", behaviour="Satır yüklenir; malzeme kapasite tüketmeden planlanır",
             message=f"Malzeme {d.NO_ROUTE_PART}'ün rotası yok ({route}); kapasite tüketmeden planlanır.",
             core_difference=[f"Routing {route}: {len(removed)} satır dosyada yok"]),
        dict(code="H2", type="Tanımsız malzeme kodu", file="SalesOrderLine.xlsx",
             rows=[row("SalesOrderLine", OrderNo="SO-1099")], key="SO-1099|1|", field="MaterialCode", value="MAM-099",
             step="Fonksiyon 04-02 referans kontrolü (geliştirme 3.4, 4.4)", severity="hata",
             behaviour="Satır yüklenmez", message="Malzeme bulunamadı: MAM-099 (partide ve geçerli veride yok).",
             core_difference=["SalesOrderLine SO-1099 yok (temiz hâlde de yok)"]),
        dict(code="H2", type="Tanımsız malzeme kodu", file="BOM.xlsx",
             rows=[row("BOM", MaterialCode="ALT-005", ComponentMaterialCode="SAT-999")], key="ALT-005|V1|SAT-999",
             field="ComponentMaterialCode", value="SAT-999", step="Fonksiyon 04-02 referans kontrolü (geliştirme 3.4, 4.4)",
             severity="hata", behaviour="Satır yüklenmez",
             message="Malzeme bulunamadı: SAT-999 (partide ve geçerli veride yok).",
             core_difference=["BOM ALT-005 → SAT-999 yok (temiz hâlde de yok)"]),
        dict(code="H3", type="Negatif miktar", file="PurchaseOrderLine.xlsx",
             rows=[row("PurchaseOrderLine", OrderNo="PO-2060", LineNo=1)], key="PO-2060|1", field="OpenQty",
             value=-50.0, step="Fonksiyon 04-03 mantıksal kontrol (geliştirme 3.4)", severity="hata",
             behaviour="Satır yüklenmez", message="OpenQty negatif olamaz (-50).",
             core_difference=["PurchaseOrderLine PO-2060|1 eksik"]),
        dict(code="H3", type="Negatif miktar", file="BOM.xlsx",
             rows=[row("BOM", MaterialCode="ALT-012", ComponentMaterialCode="SAT-010")], key="ALT-012|V1|SAT-010",
             field="Quantity", value=-2.0, step="Fonksiyon 04-03 mantıksal kontrol (geliştirme 3.4)", severity="hata",
             behaviour="Satır yüklenmez", message="Quantity negatif olamaz (-2).",
             core_difference=["BOM ALT-012 → SAT-010 eksik"]),
        dict(code="H4", type="Tekrar eden anahtar", file="Material.xlsx",
             rows=[mat_i + FIRST_DATA_ROW, mat_i + 1 + FIRST_DATA_ROW], key=d.ISOLATED_MATERIAL, field="Code",
             value=d.ISOLATED_MATERIAL, step="Fonksiyon 04-01 yapısal kontrol (geliştirme 3.4)", severity="hata",
             behaviour="İki satır da yüklenmez (hangisinin doğru olduğu bilinemez)",
             message=(f"Tekrar eden anahtar: Material.Code = {d.ISOLATED_MATERIAL} "
                      f"(satır {mat_i + FIRST_DATA_ROW} ve {mat_i + 1 + FIRST_DATA_ROW})."),
             core_difference=[f"Material {d.ISOLATED_MATERIAL} eksik (başka hiçbir dosyada geçmez)"]),
    ]
    return stg, errors


def _cell(v):
    if isinstance(v, bool):
        return 1 if v else 0
    return v


def write_excel(ds: Dataset, out: Path) -> dict:
    """Kirli hâli yazar: varlık başına Excel ve beklenen-hatalar.json. Özet döndürür."""
    out.mkdir(parents=True, exist_ok=True)
    stg, errors = make_dirty(ds)
    for entity, rows in stg.items():
        wb = Workbook()
        ws = wb.active
        ws.title = entity
        cols = stg_columns(entity)
        ws.append([c.name for c in cols])
        ws.append([c.comment or "" for c in cols])
        for r in rows:
            ws.append([_cell(r[c.name]) for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A3"
        info = wb.create_sheet("_Surum")
        for k, v in [("Set", d.SET_NAME), ("Varlık", f"stg.{entity}"), ("Tohum", ds.seed),
                     ("Üretim günü", ds.generated_on.isoformat()), ("Hafta 1", ds.start.isoformat()),
                     ("Düzen", "1. satır alan adı, 2. satır açıklama, veri 3. satırdan (geçici; 3.1 şablonu)")]:
            info.append([k, v])
        info.sheet_state = "hidden"
        wb.save(out / f"{entity}.xlsx")
    manifest = {
        "set": d.SET_NAME, "seed": ds.seed, "generated_on": ds.generated_on.isoformat(),
        "week1": ds.start.isoformat(), "first_data_row": FIRST_DATA_ROW,
        "files": {f"{e}.xlsx": len(r) for e, r in stg.items()}, "errors": errors,
    }
    (out / "beklenen-hatalar.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, default=_json_default), encoding="utf-8")
    return manifest


def _json_default(v):
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    raise TypeError(type(v))
